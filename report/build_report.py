"""Assemble the final report from report/sections/*.md and figures/*.png.

    python report/build_report.py            # builds both
    python report/build_report.py docx       # Word paper only
    python report/build_report.py pptx       # slide deck only

Outputs: report/output/ELE_IDSA_Project1_Report.docx and ELE_IDSA_Project1_Slides.pptx
Sections that are still empty and figures that don't exist yet appear as marked placeholders.

Markdown supported in the section files: #/##/### headings, "- " bullets (two-space indent = sub-bullet),
| tables |, **bold**, *italic*, `code`, and figure references. A figure is inserted where its file
name (e.g. figures/fig03_boardings_by_hour.png) is first mentioned.
"""
import re
import sys
from datetime import date
from pathlib import Path

from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SECTIONS = ROOT / "report" / "sections"
FIGURES = ROOT / "figures"
OUT = ROOT / "report" / "output"

TITLE = "Why Buses Bunch"
SUBTITLE = "The Relationship Between Passenger Demand and Bus Bunching on CapMetro Rapid Route 801, July–December 2021"
COURSE = "ELE IDSA Project 1 · GEC2B · Group 10"
INSTRUCTOR = "Instructor: Ezron John D. Cabrera"
MEMBERS = ["Badal, Khalil", "Lopez, Elijah David", "Marquez, Rence Joseph"]

# (heading level, heading text, section file or None)
OUTLINE = [
    (1, "1. Introduction", "01_introduction.md"),
    (1, "2. Methodology", "02_methodology.md"),
    (1, "3. Results and Discussion", None),
    (2, "3.1 Ridership Patterns", "03_results_ridership_patterns.md"),
    (2, "3.2 Where and When Passengers Board", "03b_results_heatmap.md"),
    (2, "3.3 Hypothesis Results", "05_results_hypotheses.md"),
    (1, "4. Conclusion", "06_conclusion.md"),
]
CAPTIONS = {
    "fig01_dwell_cleaning": "Dwell time before and after cleaning",
    "fig02_boardings_heatmap": "Average weekday boardings by stop and hour",
    "fig03_boardings_by_hour": "Boardings by hour of day, weekdays vs weekends",
    "fig04_daily_boardings": "Total boardings per day, July–December 2021",
    "fig05_headway_distribution": "Distribution of headways, weekdays 7:00–17:59",
    "fig06_h1_load_vs_headway": "H1: Passengers on board by headway group",
    "fig07_h2_dwell_vs_boardings": "H2: Dwell time vs passengers boarding",
    "fig08_h3_headway_cv_by_stop": "H3: Headway irregularity along the route",
}
FIG_RE = re.compile(r"fig\d\d_[a-z0-9_]+")
INK, MUTED, ACCENT, HEAD_FILL = "0B0B0B", "52514E", "2A78D6", "DCE8F8"


# ------------------------------------------------------------------ markdown -> blocks
def parse(md_text):
    """Return blocks: ('h', level, text) ('p', text) ('ul', [(level, text)]) ('table', rows) ('fig', name)."""
    blocks, seen = [], set()

    def figs_in(line):
        out = []
        for name in FIG_RE.findall(line):
            if name not in seen:
                seen.add(name)
                out.append(("fig", name))
        return out

    def strip_fig_refs(line):
        line = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", line)
        line = re.sub(r"`?(?:\.\./)*(?:figures/)?fig\d\d_[a-z0-9_]+(?:\.png)?`?", "", line)
        bare = re.sub(r"[*_`:·\-\s]|figure|file|see", "", line, flags=re.I)
        return "" if not bare else line

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        if not s:
            i += 1
            continue
        if s.startswith("*Owner") or s.startswith("_Owner"):
            blocks += figs_in(s)
            i += 1
            continue
        m = re.match(r"(#{1,6})\s+(.*)", s)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if re.match(r"\s*([-*+]|\d+\.)\s+", raw):
            items, pending = [], []
            while i < len(lines) and re.match(r"\s*([-*+]|\d+\.)\s+", lines[i]):
                ln = lines[i]
                level = min((len(ln) - len(ln.lstrip())) // 2, 2)
                text = re.sub(r"^\s*([-*+]|\d+\.)\s+", "", ln).strip()
                pending += figs_in(text)
                text = strip_fig_refs(text)
                if text:
                    items.append((level, text))
                i += 1
            if items:
                blocks.append(("ul", items))
            blocks += pending
            continue
        para = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"\s*(#|\||[-*+]\s|\d+\.\s)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        text = " ".join(para)
        figs = figs_in(text)
        text = strip_fig_refs(text)
        if text:
            blocks.append(("p", text))
        blocks += figs
    return blocks


def section_blocks(filename, slides=False):
    path = SECTIONS / filename
    short = path.with_name(path.stem + ".slides.md")
    if slides and short.exists():                 # optional shorter version used only for the deck
        path = short
    blocks = parse(path.read_text(encoding="utf-8")) if path.exists() else []
    if blocks and blocks[0][0] == "h" and blocks[0][1] == 1:
        blocks = blocks[1:]                       # the file's own title is replaced by the outline heading
    content = [b for b in blocks if b[0] != "h"]
    if not content:
        owner = re.search(r"Owner:\s*([A-Za-z]+)", path.read_text(encoding="utf-8")) if path.exists() else None
        return [("pending", f"[Section pending{': ' + owner.group(1) if owner else ''}]")]
    # demote remaining headings so they sit under the outline heading
    return [("h", b[1] + 1, b[2]) if b[0] == "h" else b for b in blocks]


def inline_runs(text):
    """Split **bold**, *italic* and `code` into (text, bold, italic) runs."""
    runs = []
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            runs.append((part[2:-2], True, False))
        elif part.startswith("*") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        elif part.startswith("`"):
            runs.append((part[1:-1], False, False))
        else:
            runs.append((part, False, False))
    return runs


def plain(text):
    return "".join(t for t, _, _ in inline_runs(text))


def all_content(slides=False):
    """[(level, heading, blocks)] in report order, plus the figure numbering."""
    out, numbers = [], {}
    for level, heading, fname in OUTLINE:
        blocks = section_blocks(fname, slides) if fname else []
        for b in blocks:
            if b[0] == "fig" and b[1] not in numbers:
                numbers[b[1]] = len(numbers) + 1
        out.append((level, heading, blocks))
    return out, numbers


def caption(name, numbers):
    return f"Figure {numbers[name]}. {CAPTIONS.get(name, name.split('_', 1)[1].replace('_', ' ').capitalize())}"


# ------------------------------------------------------------------ Word
def build_docx(content, numbers):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12), ("Heading 4", 12)]:
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(size), True
        st.font.color.rgb = RGBColor.from_string(INK)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
        st.paragraph_format.keep_with_next = True

    def add_runs(p, text, size=None, color=None):
        for t, b, it in inline_runs(text):
            r = p.add_run(t)
            r.bold, r.italic = b or None, it or None
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def centered(text, size, bold=False, after=6, color=INK):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)

    # title page
    for _ in range(6):
        doc.add_paragraph()
    centered(TITLE, 24, bold=True, after=8)
    centered(SUBTITLE, 14, after=48)
    centered(COURSE, 12)
    centered(INSTRUCTOR, 12, after=36)
    centered("Prepared by", 12, color=MUTED)
    for m in MEMBERS:
        centered(m, 12, after=2)
    centered(date.today().strftime("%B %Y"), 12, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    for level, heading, blocks in content:
        doc.add_heading(heading, level=level)
        for b in blocks:
            kind = b[0]
            if kind == "h":
                doc.add_heading(plain(b[2]), level=min(b[1], 4))
            elif kind == "p":
                p = doc.add_paragraph()
                add_runs(p, b[1])
                nxt = blocks[blocks.index(b) + 1] if blocks.index(b) + 1 < len(blocks) else None
                p.paragraph_format.keep_with_next = bool(nxt and nxt[0] in ("table", "fig"))
            elif kind == "pending":
                add_runs(doc.add_paragraph(), b[1], color=ACCENT)
            elif kind == "ul":
                for lvl, text in b[1]:
                    p = doc.add_paragraph(style="List Bullet" if lvl == 0 else "List Bullet 2")
                    p.paragraph_format.space_after = Pt(3)
                    add_runs(p, text)
            elif kind == "table":
                rows = b[1]
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(0)
                        add_runs(p, row[ci] if ci < len(row) else "", size=10.5)
                        p.paragraph_format.keep_with_next = ri < len(rows) - 1   # keep the table on one page
                        cell.width = Cm(16 / ncol)
                        if ri == 0:
                            for r in p.runs:
                                r.bold = True
                            shade(cell, HEAD_FILL)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif kind == "fig":
                name = b[1]
                img = FIGURES / f"{name}.png"
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                if img.exists():
                    w, h = Image.open(img).size
                    width = Cm(16)
                    if h / w * 16 > 18:            # keep tall figures on one page
                        width = Cm(18 * w / h)
                    p.add_run().add_picture(str(img), width=width)
                else:
                    add_runs(p, f"[Figure pending: {name}.png]", color=ACCENT)
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(cap, caption(name, numbers), size=10.5, color=MUTED)

    # page numbers in the footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    sec.different_first_page_header_footer = True

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "ELE_IDSA_Project1_Report.docx"
    doc.save(path)
    return path


# ------------------------------------------------------------------ Slides
def words(items):
    return sum(len(plain(t).split()) for _, t in items)


def build_pptx(content, numbers):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H, M = 13.333, 7.5, 0.6
    FONT = "Calibri"
    rgb = RGBColor.from_string

    def text_box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        return tf

    def put_runs(par, text, size, color=INK, bold=None):
        for t, b, it in inline_runs(text):
            r = par.add_run()
            r.text = t
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = bold if bold is not None else (b or None)
            r.font.italic = it or None
            r.font.color.rgb = rgb(color)

    def new_slide(title, kicker=None):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(ACCENT); bar.line.fill.background()
        if kicker:
            put_runs(text_box(s, M, 0.32, W - 2 * M, 0.35).paragraphs[0], kicker.upper(), 12, ACCENT, bold=True)
        put_runs(text_box(s, M, 0.62, W - 2 * M, 0.8).paragraphs[0], title, 28, INK, bold=True)
        foot = text_box(s, M, H - 0.45, W - 2 * M, 0.3)
        put_runs(foot.paragraphs[0], COURSE, 10, MUTED)
        num = text_box(s, W - M - 1, H - 0.45, 1, 0.3).paragraphs[0]
        num.alignment = PP_ALIGN.RIGHT
        put_runs(num, str(len(prs.slides)), 10, MUTED)
        return s

    def bullets(slide, items, x, y, w, h, size):
        tf = text_box(slide, x, y, w, h)
        first = True
        for lvl, text in items:
            par = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            par.space_after = Pt(8)
            par.level = lvl
            put_runs(par, ("•  " if lvl == 0 else "–  ") + text, size - 2 * lvl)

    def chunk(items, budget):
        out, cur = [], []
        for it in items:
            if cur and words(cur + [it]) > budget:
                out.append(cur); cur = []
            cur.append(it)
        if cur:
            out.append(cur)
        return out

    # title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb("0D366B"); bg.line.fill.background()
    put_runs(text_box(s, 0.9, 1.7, W - 1.8, 1.2).paragraphs[0], TITLE, 48, "FFFFFF", bold=True)
    put_runs(text_box(s, 0.9, 2.85, W - 2.6, 1.3).paragraphs[0], SUBTITLE, 22, "CDE2FB")
    tf = text_box(s, 0.9, 4.7, W - 1.8, 2)
    put_runs(tf.paragraphs[0], COURSE + "  ·  " + INSTRUCTOR, 16, "FFFFFF")
    put_runs(tf.add_paragraph(), ", ".join(MEMBERS), 16, "9EC5F4")

    for level, heading, blocks in content:
        if not blocks:
            continue                                   # "3. Results and Discussion" is only a parent heading
        title = re.sub(r"^\d+(\.\d+)*\.?\s+", "", heading)
        kicker = "Results and Discussion" if heading.startswith("3.") else title
        sub = None
        pending_text = []                              # bullets/paragraphs waiting to be placed

        def flush_text(items, slide_title):
            for part in chunk(items, 95):
                s = new_slide(slide_title, kicker)
                bullets(s, part, M, 1.65, W - 2 * M, H - 2.4, 20)

        i = 0
        while i < len(blocks):
            b = blocks[i]
            kind = b[0]
            if kind == "h":
                carry = (len(pending_text) == 1 and words(pending_text) <= 30
                         and i + 1 < len(blocks) and blocks[i + 1][0] == "fig")   # short line rides onto the next figure slide
                if pending_text and not carry:
                    flush_text(pending_text, sub or title); pending_text = []
                sub = plain(b[2])
            elif kind in ("p", "pending"):
                pending_text.append((0, b[1]))
            elif kind == "ul":
                pending_text += b[1]
            elif kind == "table":
                intro = None
                if len(pending_text) == 1 and words(pending_text) <= 35:
                    intro = pending_text[0][1]; pending_text = []
                if pending_text:
                    flush_text(pending_text, sub or title); pending_text = []
                rows = b[1]
                s = new_slide(sub or title, kicker)
                top = 1.7
                if intro:
                    put_runs(text_box(s, M, 1.55, W - 2 * M, 0.5).paragraphs[0], intro, 16, MUTED)
                    top = 2.2
                ncol = max(len(r) for r in rows)
                avail = H - top - 0.8
                rh = min(0.5, avail / len(rows))
                shape = s.shapes.add_table(len(rows), ncol, Inches(M), Inches(top), Inches(W - 2 * M), Inches(rh * len(rows)))
                tbl = shape.table
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = tbl.cell(ri, ci)
                        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(HEAD_FILL if ri == 0 else "FFFFFF")
                        par = cell.text_frame.paragraphs[0]
                        put_runs(par, row[ci] if ci < len(row) else "", 14 if len(rows) < 9 else 12,
                                 INK, bold=True if ri == 0 else None)
            elif kind == "fig":
                name = b[1]
                # bullets right after the figure (until the next figure/heading/table) go beside it
                side = []
                j = i + 1
                while j < len(blocks) and blocks[j][0] in ("p", "ul"):
                    side += [(0, blocks[j][1])] if blocks[j][0] == "p" else blocks[j][1]
                    j += 1
                intro = None
                if len(pending_text) == 1 and words(pending_text) <= 30:   # one short line: put it on the figure slide
                    intro = pending_text[0][1]; pending_text = []
                if pending_text:
                    flush_text(pending_text, sub or title); pending_text = []
                s = new_slide(caption(name, numbers).split(". ", 1)[1], kicker)
                img = FIGURES / f"{name}.png"
                fig_w = 8.2 if side else W - 2 * M
                box_w, box_h, top = fig_w, H - 2.55, 1.6
                if intro:
                    put_runs(text_box(s, M, 1.4, W - 2 * M, 0.4).paragraphs[0], intro, 14, MUTED)
                    box_h, top = box_h - 0.35, 1.95
                if img.exists():
                    w, h = Image.open(img).size
                    scale = min(box_w / w, box_h / h)
                    s.shapes.add_picture(str(img), Inches(M), Inches(top), Inches(w * scale), Inches(h * scale))
                    cap_y = top + h * scale + 0.05
                else:
                    ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M), Inches(top), Inches(box_w), Inches(box_h - 0.4))
                    ph.fill.background(); ph.line.color.rgb = rgb(ACCENT); ph.line.dash_style = 4
                    put_runs(ph.text_frame.paragraphs[0], f"[Figure pending: {name}.png]", 18, ACCENT)
                    cap_y = top + box_h - 0.35
                put_runs(text_box(s, M, cap_y, fig_w, 0.3).paragraphs[0], caption(name, numbers), 11, MUTED)
                if side:
                    extra = []
                    if words(side) > 85:               # too much for the side panel: overflow to a text slide
                        side, extra = chunk(side, 85)[0], [it for c in chunk(side, 85)[1:] for it in c]
                    size = 16 if words(side) < 60 else 14
                    bullets(s, side, M + fig_w + 0.35, top + 0.1, W - 2 * M - fig_w - 0.35, H - top - 0.9, size)
                    if extra:
                        flush_text(extra, caption(name, numbers).split(". ", 1)[1] + " (cont.)")
                i = j
                continue
            i += 1
        if pending_text:
            flush_text(pending_text, sub or title)

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "ELE_IDSA_Project1_Slides.pptx"
    prs.save(path)
    return path


def main():
    which = sys.argv[1] if len(sys.argv) > 1 else "both"
    if which in ("both", "docx"):
        print("wrote", build_docx(*all_content()))
    if which in ("both", "pptx"):
        print("wrote", build_pptx(*all_content(slides=True)))
    missing = [f for f in CAPTIONS if not (FIGURES / f"{f}.png").exists()]
    if missing:
        print("figures still missing:", ", ".join(missing))


if __name__ == "__main__":
    main()
